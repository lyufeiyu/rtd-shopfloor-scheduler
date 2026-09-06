import sys
import re
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import matplotlib.dates as mdates
from openpyxl.utils import get_column_letter
plt.rcParams["font.sans-serif"] = ["SimHei"]
plt.rcParams["axes.unicode_minus"] = False

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")


def load_data(excel_path):
    # =========================================================
    # 1. 待读取 Job
    # =========================================================
    pending_df = pd.read_excel(excel_path, sheet_name="8-制造单批次流转表")

    for col in ["制造单号", "工艺路线", "已完成工站", "下一工站"]:
        pending_df[col] = pending_df[col].apply(
            lambda x: x.strip() if isinstance(x, str) else x
        )

    pending_df["数量"] = pd.to_numeric(pending_df["数量"], errors="coerce")
    pending_df = pending_df.dropna(
        subset=["制造单号", "批次", "数量", "工艺路线", "下一工站"]
    )


    # =========================================================
    # 2. 工艺路线 + 工站 -> 允许设备
    # =========================================================
    route_df = pd.read_excel(excel_path, sheet_name="5-工艺路线-设备表")

    # 处理 Excel 合并单元格
    route_df["工艺路线"] = route_df["工艺路线"].ffill()

    for col in ["工艺路线", "工站"]:
        route_df[col] = route_df[col].apply(
            lambda x: x.strip() if isinstance(x, str) else x
        )

    route_equipment_map = {}

    for _, row in route_df.iterrows():
        if pd.isna(row["工艺路线"]) or pd.isna(row["工站"]) or pd.isna(row["设备"]):
            continue

        key = (row["工艺路线"], row["工站"])

        machines = re.split(r"[,，;；\r\n]+", str(row["设备"]))
        machines = [m.strip() for m in machines if m.strip()]

        route_equipment_map.setdefault(key, [])

        for machine in machines:
            if machine not in route_equipment_map[key]:
                route_equipment_map[key].append(machine)


    # =========================================================
    # 3. 工站 -> 所有机器
    # =========================================================
    station_df = pd.read_excel(excel_path, sheet_name="工站-设备表")

    for col in ["设备名", "工站"]:
        station_df[col] = station_df[col].apply(
            lambda x: x.strip() if isinstance(x, str) else x
        )

    station_df = station_df.dropna(subset=["设备名", "工站"])

    station_equipment_map = (
        station_df.groupby("工站")["设备名"]
        .apply(list)
        .to_dict()
    )


    # =========================================================
    # 4. 设备 -> UPH
    # =========================================================
    machine_df = pd.read_excel(excel_path, sheet_name="3-设备表")

    machine_df["设备名"] = machine_df["设备名"].apply(
        lambda x: x.strip() if isinstance(x, str) else x
    )

    machine_df["设备产能UPH"] = pd.to_numeric(
        machine_df["设备产能UPH"],
        errors="coerce"
    )

    machine_df = machine_df.dropna(subset=["设备名", "设备产能UPH"])

    machine_uph = dict(
        zip(machine_df["设备名"], machine_df["设备产能UPH"])
    )


    # =========================================================
    # 5. 当前 WIP 快照
    # =========================================================
    wip_df = pd.read_excel(excel_path, sheet_name="6-WIP在制表")

    for col in ["设备", "工艺路线"]:
        wip_df[col] = wip_df[col].apply(
            lambda x: x.strip() if isinstance(x, str) else x
        )

    wip_df["投产时间"] = pd.to_datetime(
        wip_df["投产时间"],
        errors="coerce"
    )

    wip_df["数量"] = pd.to_numeric(
        wip_df["数量"],
        errors="coerce"
    )

    wip_df = wip_df.dropna(
        subset=["设备", "工艺路线", "投产时间", "数量"]
    )

    return pending_df, route_equipment_map, station_equipment_map, machine_uph, wip_df


def schedule_station(
    pending_df,
    route_equipment_map,
    station_equipment_map,
    machine_uph,
    wip_df,
    station,
    current_time
):
    current_time = pd.Timestamp(current_time)

    # =========================================================
    # 1. 当前工站所有待排 Job
    # =========================================================
    jobs = pending_df[
        pending_df["下一工站"] == station
    ].copy().reset_index(drop=True)

    if jobs.empty:
        print(f"{station} 没有待排 Job")
        return None

    # Job 唯一名称
    jobs["job_id"] = [
        f"{row['制造单号']}-{row['批次']}"
        for _, row in jobs.iterrows()
    ]


    # =========================================================
    # 2. 当前工站所有机器
    # =========================================================
    machines = station_equipment_map.get(station, [])

    # 必须存在有效 UPH
    machines = [
        m for m in machines
        if m in machine_uph and machine_uph[m] > 0
    ]

    if not machines:
        print(f"{station} 没有有效机器")
        return None

    machine_index = {
        machine: j for j, machine in enumerate(machines)
    }


    # =========================================================
    # 3. 基础可加工矩阵 + 加工时间矩阵
    #
    # base_assignment:
    # 1 = 工艺路线理论允许使用这台机器
    # 0 = 不允许
    #
    # processing_time:
    # 数量 / UPH * 60
    # 单位：分钟
    # =========================================================
    base_assignment = np.zeros(
        (len(jobs), len(machines)),
        dtype=int
    )

    processing_time = np.full(
        (len(jobs), len(machines)),
        np.inf
    )

    for i, job in jobs.iterrows():
        route = job["工艺路线"]
        quantity = job["数量"]

        allowed_machines = route_equipment_map.get(
            (route, station),
            []
        )

        for machine in allowed_machines:
            if machine not in machine_index:
                continue

            j = machine_index[machine]

            base_assignment[i, j] = 1
            processing_time[i, j] = quantity / machine_uph[machine] * 60


    # =========================================================
    # 4. 根据 WIP 快照计算机器初始释放时间
    #
    # machine_busy_until:
    # machine -> 什么时候可以重新使用
    #
    # 如果没有 WIP，则等于 current_time
    # =========================================================
    machine_busy_until = {
        machine: current_time
        for machine in machines
    }

    initial_wip = []

    for _, wip in wip_df.iterrows():
        machine = wip["设备"]

        if machine not in machine_index:
            continue

        uph = machine_uph.get(machine)

        if uph is None or uph <= 0:
            raise ValueError(
                f"WIP设备 {machine} 没有有效UPH，无法计算释放时间"
            )

        start_time = wip["投产时间"]
        process_minutes = wip["数量"] / uph * 60

        finish_time = start_time + pd.Timedelta(
            minutes=process_minutes
        )

        # 只有当前时间还没完成的 WIP 才算占用
        if start_time <= current_time < finish_time:
            if finish_time > machine_busy_until[machine]:
                machine_busy_until[machine] = finish_time

            initial_wip.append({
                "machine": machine,
                "route": wip["工艺路线"],
                "quantity": wip["数量"],
                "start_time": start_time,
                "finish_time": finish_time
            })


    # =========================================================
    # 5. 初始动态 Job-Machine 可分配矩阵
    #
    # base_assignment = 1
    # 并且机器当前空闲
    # 才能为 1
    # =========================================================
    initial_assignment = base_assignment.copy()

    for j, machine in enumerate(machines):
        if machine_busy_until[machine] > current_time:
            initial_assignment[:, j] = 0


    job_names = jobs["job_id"].tolist()

    initial_assignment_df = pd.DataFrame(
        initial_assignment,
        index=job_names,
        columns=machines
    )

    processing_time_df = pd.DataFrame(
        processing_time,
        index=job_names,
        columns=machines
    )


    # =========================================================
    # 6. 启发式调度
    #
    # 规则：
    # ① 可选机器最少的 Job 优先
    # ② 同等情况下，加工时间短的优先
    # ③ 为 Job 选择加工时间最短的空闲机器
    #
    # 当没有任务可以立即派时：
    # current_time 跳到最近一个机器完成时间
    # =========================================================
    unscheduled = set(range(len(jobs)))
    schedule = []

    while unscheduled:

        # -----------------------------------------------------
        # 当前时刻持续派工，直到没有可派组合
        # -----------------------------------------------------
        while True:
            candidates = []

            for i in unscheduled:
                available_machines = []

                for j, machine in enumerate(machines):
                    if (
                        base_assignment[i, j] == 1
                        and machine_busy_until[machine] <= current_time
                    ):
                        available_machines.append(j)

                if not available_machines:
                    continue

                # 当前 Job 有多少台机器可以立即使用
                flexibility = len(available_machines)

                # 给当前 Job 选择加工时间最短的机器
                best_machine_index = min(
                    available_machines,
                    key=lambda j: processing_time[i, j]
                )

                best_process_time = processing_time[
                    i,
                    best_machine_index
                ]

                candidates.append(
                    (
                        flexibility,
                        best_process_time,
                        i,
                        best_machine_index
                    )
                )

            # 当前时间没有任何 Job 可以派
            if not candidates:
                break

            # 可选机器少的 Job 优先
            # 相同时加工时间短的优先
            _, process_minutes, job_index, machine_idx = min(
                candidates,
                key=lambda x: (x[0], x[1])
            )

            job = jobs.iloc[job_index]
            machine = machines[machine_idx]

            start_time = current_time
            finish_time = start_time + pd.Timedelta(
                minutes=float(process_minutes)
            )

            schedule.append({
                "job_id": job["job_id"],
                "制造单号": job["制造单号"],
                "批次": job["批次"],
                "工艺路线": job["工艺路线"],
                "工站": station,
                "数量": job["数量"],
                "设备": machine,
                "UPH": machine_uph[machine],
                "开始时间": start_time,
                "完成时间": finish_time,
                "加工时间(分钟)": round(float(process_minutes), 2)
            })

            # 派工后机器立即被占用
            machine_busy_until[machine] = finish_time

            # Job 已经完成派工
            unscheduled.remove(job_index)


        # 所有 Job 都派完
        if not unscheduled:
            break


        # =====================================================
        # 7. 当前没有 Job 可以派
        #
        # 找与剩余 Job 有关的机器中，
        # 最近一个完成时间。
        #
        # 时间直接跳过去。
        # =====================================================
        next_events = []

        for j, machine in enumerate(machines):

            if machine_busy_until[machine] <= current_time:
                continue

            # 判断这台机器是否能加工某个剩余 Job
            needed = any(
                base_assignment[i, j] == 1
                for i in unscheduled
            )

            if needed:
                next_events.append(
                    machine_busy_until[machine]
                )

        # 没有未来机器可以释放给剩余 Job
        # 说明这些 Job 无法在当前工站找到可加工设备
        if not next_events:
            break

        # 事件驱动，不一分钟一分钟模拟
        current_time = min(next_events)


    # =========================================================
    # 8. 调度结果
    # =========================================================
    schedule_df = pd.DataFrame(schedule)

    unscheduled_df = jobs.iloc[
        sorted(unscheduled)
    ].copy()

    return {
        "jobs": jobs,
        "machines": machines,
        "initial_wip": pd.DataFrame(initial_wip),
        "initial_assignment_matrix": initial_assignment_df,
        "processing_time_matrix": processing_time_df,
        "schedule": schedule_df,
        "unscheduled_jobs": unscheduled_df,
        "machine_busy_until": machine_busy_until
    }





def plot_gantt(result, current_time):
    schedule = result["schedule"]
    initial_wip = result["initial_wip"]

    current_time = pd.Timestamp(current_time)

    # ---------------------------------------------------------
    # 只画真正有 WIP 或有排产任务的机器
    # ---------------------------------------------------------
    used_machines = set()

    if not initial_wip.empty:
        used_machines.update(initial_wip["machine"].tolist())

    if not schedule.empty:
        used_machines.update(schedule["设备"].tolist())

    used_machines = sorted(used_machines)

    if not used_machines:
        print("没有可绘制的排产结果")
        return

    y_pos = {machine: i for i, machine in enumerate(used_machines)}

    # 机器越多，图自动变高
    fig_height = max(6, len(used_machines) * 0.38)

    fig, ax = plt.subplots(figsize=(18, fig_height))

    # ---------------------------------------------------------
    # 1. 初始 WIP
    #
    # 关键：
    # 甘特图从 current_time 开始看，
    # 不再把之前已经生产的那一大段画出来
    # ---------------------------------------------------------
    if not initial_wip.empty:
        for _, row in initial_wip.iterrows():
            machine = row["machine"]

            if machine not in y_pos:
                continue

            start = max(
                pd.Timestamp(row["start_time"]),
                current_time
            )

            finish = pd.Timestamp(row["finish_time"])

            if finish <= current_time:
                continue

            ax.barh(
                y_pos[machine],
                finish - start,
                left=start,
                height=0.65,
                alpha=0.35,
                hatch="//"
            )

            # WIP 只写一个简短标记
            middle = start + (finish - start) / 2

            ax.text(
                middle,
                y_pos[machine],
                "WIP",
                ha="center",
                va="center",
                fontsize=7
            )

    # ---------------------------------------------------------
    # 2. 启发式算法新安排的 Job
    # ---------------------------------------------------------
    if not schedule.empty:
        for _, row in schedule.iterrows():
            machine = row["设备"]

            if machine not in y_pos:
                continue

            start = pd.Timestamp(row["开始时间"])
            finish = pd.Timestamp(row["完成时间"])

            ax.barh(
                y_pos[machine],
                finish - start,
                left=start,
                height=0.65
            )

            # 只有条足够长才显示 Job 名
            duration_minutes = (
                finish - start
            ).total_seconds() / 60

            if duration_minutes >= 60:
                job_label = str(row["job_id"])

                middle = start + (finish - start) / 2

                ax.text(
                    middle,
                    y_pos[machine],
                    job_label,
                    ha="center",
                    va="center",
                    fontsize=6
                )

    # ---------------------------------------------------------
    # 3. 当前时刻参考线
    # ---------------------------------------------------------
    ax.axvline(
        current_time,
        linestyle="--",
        linewidth=1.5,
        label="排产开始时间"
    )

    # ---------------------------------------------------------
    # 4. Y轴
    # ---------------------------------------------------------
    ax.set_yticks(range(len(used_machines)))
    ax.set_yticklabels(used_machines, fontsize=8)

    # 第一台机器显示在顶部
    ax.invert_yaxis()

    # ---------------------------------------------------------
    # 5. X轴
    # ---------------------------------------------------------
    ax.xaxis.set_major_locator(
        mdates.AutoDateLocator()
    )

    ax.xaxis.set_major_formatter(
        mdates.DateFormatter("%m-%d\n%H:%M")
    )

    # ---------------------------------------------------------
    # 6. 自动限制显示时间范围
    # ---------------------------------------------------------
    all_finish_times = []

    if not initial_wip.empty:
        all_finish_times.extend(
            pd.to_datetime(initial_wip["finish_time"]).tolist()
        )

    if not schedule.empty:
        all_finish_times.extend(
            pd.to_datetime(schedule["完成时间"]).tolist()
        )

    if all_finish_times:
        max_finish = max(all_finish_times)

        ax.set_xlim(
            current_time - pd.Timedelta(minutes=20),
            max_finish + pd.Timedelta(minutes=30)
        )

    # ---------------------------------------------------------
    # 7. 图形样式
    # ---------------------------------------------------------
    ax.set_xlabel("时间")
    ax.set_ylabel("设备")
    ax.set_title("单工站排产甘特图")

    ax.grid(
        axis="x",
        linestyle="--",
        alpha=0.25
    )

    ax.legend(loc="upper right")

    plt.tight_layout()
    plt.show()


def export_schedule(result, station, output_path=None):
    if output_path is None:
        output_path = f"{station}_排产方案.xlsx"

    schedule_df = result["schedule"].copy()
    initial_wip_df = result["initial_wip"].copy()
    unscheduled_df = result["unscheduled_jobs"].copy()

    # 排产方案
    if not schedule_df.empty:
        schedule_df = schedule_df[
            [
                "job_id", "制造单号", "批次", "工艺路线", "工站",
                "数量", "设备", "UPH", "开始时间", "完成时间",
                "加工时间(分钟)"
            ]
        ]

        # pandas Timestamp -> Python datetime
        for col in ["开始时间", "完成时间"]:
            schedule_df[col] = pd.Series(
                [
                    x.to_pydatetime() if isinstance(x, pd.Timestamp) else x
                    for x in schedule_df[col]
                ],
                index=schedule_df.index,
                dtype=object
            )

    # 初始 WIP
    if not initial_wip_df.empty:
        initial_wip_df = initial_wip_df.rename(columns={
            "machine": "设备",
            "route": "工艺路线",
            "quantity": "数量",
            "start_time": "投产时间",
            "finish_time": "预计完成时间"
        })

        for col in ["投产时间", "预计完成时间"]:
            if col in initial_wip_df.columns:
                initial_wip_df[col] = pd.Series(
                    [
                        x.to_pydatetime() if isinstance(x, pd.Timestamp) else x
                        for x in initial_wip_df[col]
                    ],
                    index=initial_wip_df.index,
                    dtype=object
                )

    # 未排产任务
    if not unscheduled_df.empty:
        keep_cols = [
            col for col in
            ["job_id", "制造单号", "批次", "工艺路线", "下一工站", "数量"]
            if col in unscheduled_df.columns
        ]
        unscheduled_df = unscheduled_df[keep_cols]

    # 写 Excel
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        schedule_df.to_excel(writer, sheet_name="排产方案", index=False)
        initial_wip_df.to_excel(writer, sheet_name="初始WIP", index=False)
        unscheduled_df.to_excel(writer, sheet_name="未排产任务", index=False)

        for sheet_name in writer.book.sheetnames:
            ws = writer.book[sheet_name]

            # 冻结首行
            ws.freeze_panes = "A2"

            # 筛选
            if ws.max_row > 1:
                ws.auto_filter.ref = ws.dimensions

            # 获取表头
            headers = {
                ws.cell(1, col).value: col
                for col in range(1, ws.max_column + 1)
            }

            # 自动列宽
            for col in range(1, ws.max_column + 1):
                max_len = 0

                for row in range(1, ws.max_row + 1):
                    value = ws.cell(row, col).value
                    if value is not None:
                        max_len = max(max_len, len(str(value)))

                width = min(max(max_len + 2, 10), 30)
                ws.column_dimensions[get_column_letter(col)].width = width

            # 时间格式 + 加宽
            for name in ["开始时间", "完成时间", "投产时间", "预计完成时间"]:
                if name in headers:
                    col = headers[name]
                    ws.column_dimensions[get_column_letter(col)].width = 21

                    for row in range(2, ws.max_row + 1):
                        ws.cell(row, col).number_format = "yyyy-mm-dd hh:mm:ss"

            # 加工时间保留两位小数
            if "加工时间(分钟)" in headers:
                col = headers["加工时间(分钟)"]

                for row in range(2, ws.max_row + 1):
                    ws.cell(row, col).number_format = "0.00"

    print(f"排产方案已输出: {output_path}")





if __name__ == "__main__":
    excel_path = "RTD_Dataset_v6.xlsx"

    pending_df, route_equipment_map, station_equipment_map, machine_uph, wip_df = (
        load_data(excel_path)
    )

    # =========================================================
    # 只排一个工站
    # =========================================================
    station = "DB1"

    # 这里填写 WIP 快照对应的时刻
    current_time = "2026-08-12 10:00:00"

    result = schedule_station(
        pending_df,
        route_equipment_map,
        station_equipment_map,
        machine_uph,
        wip_df,
        station,
        current_time
    )

    if result is not None:

        print("\n==============================")
        print("当前工站:", station)
        print("==============================")

        print("\n初始 WIP 占用:")
        if result["initial_wip"].empty:
            print("无")
        else:
            print(result["initial_wip"])

        print("\n初始 Job-Machine 可分配矩阵:")
        print(result["initial_assignment_matrix"])

        print("\nJob-Machine 加工时间矩阵（分钟）:")
        print(result["processing_time_matrix"].round(2))

        print("\n最终排产结果:")
        print(result["schedule"])

        print("\n未能排产的 Job:")
        if result["unscheduled_jobs"].empty:
            print("无")
        else:
            print(
                result["unscheduled_jobs"][
                    ["job_id", "工艺路线", "下一工站", "数量"]
                ]
            )

        if result is not None:
            print("\n最终排产结果:")
            print(result["schedule"])

            # 输出 Excel
            export_schedule(result, station)

            # 画甘特图
            plot_gantt(result, current_time)